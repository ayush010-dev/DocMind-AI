from openpyxl import load_workbook

def extract_text(file_path):
    workbook = load_workbook(file_path)
    blocks = []

    for sheet in workbook.worksheets:
        text = f"Sheet: {sheet.title}\n"
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join(
                str(cell) for cell in row if cell is not None
            )
            if row_text.strip():
                text += row_text + "\n"
                
        if text.strip() != f"Sheet: {sheet.title}":
            blocks.append({
                "text": text,
                "sheet": sheet.title
            })

    return blocks